"""
import_docs.py
----------------
Extraccion de texto de documentos Word (.docx) y PDF (.pdf) para importar
el menu. El texto extraido se le pasa al LLM (Ollama) para que identifique
los platos; por eso Word/PDF de formato libre funcionan, aunque con menos
fiabilidad que un Excel estructurado.

Requiere (instalar en el entorno):
    pip install python-docx pypdf
"""
import io


def disponible_docx() -> bool:
    try:
        import docx  # python-docx
        return True
    except Exception:
        return False


def disponible_pdf() -> bool:
    try:
        import pypdf
        return True
    except Exception:
        return False


def extraer_texto_docx(datos: bytes) -> str:
    import docx
    documento = docx.Document(io.BytesIO(datos))
    partes = [p.text for p in documento.paragraphs if p.text and p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text and c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return "\n".join(partes)


def extraer_texto_pdf(datos: bytes) -> str:
    import pypdf
    lector = pypdf.PdfReader(io.BytesIO(datos))
    return "\n".join((pagina.extract_text() or "") for pagina in lector.pages)
