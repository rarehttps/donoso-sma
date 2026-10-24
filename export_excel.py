"""
export_excel.py
----------------
Genera archivos Excel (.xlsx) reales a partir de los datos del sistema,
con encabezados con el color de DonOso, anchos automaticos y fila fija.

Requiere openpyxl:
    pip install openpyxl

Se usa desde app_principal.py para exportar "las cosas importantes"
(pedidos/SLA, anulados, estres, inventario) en un solo libro, o los
anulados por separado desde el panel de Caja.
"""
import io

AMBAR = "B45309"   # amber-700, encabezado DonOso


def openpyxl_disponible() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def construir_libro(hojas: list) -> bytes:
    """
    hojas: lista de (nombre_hoja, columnas, filas) donde
      - columnas: lista de (clave, titulo)
      - filas: lista de dicts
    Devuelve los bytes del .xlsx.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)  # quita la hoja vacia por defecto

    relleno = PatternFill("solid", fgColor=AMBAR)
    fuente = Font(bold=True, color="FFFFFF")
    centro = Alignment(horizontal="center", vertical="center")

    for nombre, columnas, filas in hojas:
        ws = wb.create_sheet(title=str(nombre)[:31])
        ws.append([c[1] for c in columnas])
        for celda in ws[1]:
            celda.fill = relleno
            celda.font = fuente
            celda.alignment = centro
        for fila in filas:
            ws.append([_valor(fila.get(c[0])) for c in columnas])
        # ancho automatico (acotado) por columna
        for i, (clave, titulo) in enumerate(columnas, start=1):
            ancho = max([len(str(titulo))] + [len(str(_valor(f.get(clave)))) for f in filas] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 48)
        ws.freeze_panes = "A2"

    buff = io.BytesIO()
    wb.save(buff)
    return buff.getvalue()


def _valor(v):
    return "" if v is None else v
